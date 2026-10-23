# -*- coding: utf-8 -*-
"""Generate SYDT Chat end-user product documentation as a Word (.docx) file."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT = ROOT / "SYDT_Chat_Urun_Tanitim_Dokumani.docx"

# Visual theme (cyber ops — not purple-default)
BG = (18, 24, 32)
BOX = (30, 42, 56)
BOX_ACCENT = (14, 165, 160)  # teal
BOX_ALT = (45, 55, 72)
TEXT = (236, 242, 248)
MUTED = (148, 163, 184)
ARROW = (94, 234, 212)
TITLE_C = (45, 212, 191)


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "C:/Windows/Fonts/segoeui.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibri.ttf",
    ]
    if bold:
        candidates = [
            "C:/Windows/Fonts/segoeuib.ttf",
            "C:/Windows/Fonts/arialbd.ttf",
            "C:/Windows/Fonts/calibrib.ttf",
        ] + candidates
    for path in candidates:
        try:
            return ImageFont.truetype(path, size)
        except OSError:
            continue
    return ImageFont.load_default()


def rounded_rect(draw: ImageDraw.ImageDraw, xy, fill, radius=12, outline=None, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def box(draw, x, y, w, h, lines: list[str], accent=False):
    fill = BOX if not accent else (20, 60, 58)
    outline = BOX_ACCENT if accent else (71, 85, 105)
    rounded_rect(draw, (x, y, x + w, y + h), fill=fill, outline=outline, width=2)
    f = font(15, bold=True)
    fm = font(12)
    cy = y + 10
    for i, line in enumerate(lines):
        use = f if i == 0 else fm
        color = TITLE_C if i == 0 and accent else TEXT
        draw.text((x + 12, cy), line, font=use, fill=color)
        cy += 20 if i == 0 else 16


def arrow(draw, x1, y1, x2, y2):
    draw.line((x1, y1, x2, y2), fill=ARROW, width=3)
    # simple arrow head toward (x2,y2)
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        draw.polygon(
            [(x2, y2), (x2 - 10 * direction, y2 - 6), (x2 - 10 * direction, y2 + 6)],
            fill=ARROW,
        )
    else:
        direction = 1 if y2 > y1 else -1
        draw.polygon(
            [(x2, y2), (x2 - 6, y2 - 10 * direction), (x2 + 6, y2 - 10 * direction)],
            fill=ARROW,
        )


def title_bar(draw, w, title: str, subtitle: str):
    draw.rectangle((0, 0, w, 56), fill=(12, 18, 26))
    draw.text((24, 10), title, font=font(20, True), fill=TITLE_C)
    draw.text((24, 34), subtitle, font=font(12), fill=MUTED)


def save_img(img: Image.Image, name: str) -> Path:
    ASSETS.mkdir(parents=True, exist_ok=True)
    path = ASSETS / name
    img.save(path, "PNG")
    return path


def chart_architecture() -> Path:
    w, h = 1100, 620
    img = Image.new("RGB", (w, h), BG)
    d = ImageDraw.Draw(img)
    title_bar(d, w, "SYDT Chat — Sistem Mimarisi", "Son kullanıcı deneyimini taşıyan katmanlar")
    box(d, 60, 100, 280, 120, ["Web Uygulaması", "Next.js + React", "Tarayıcı arayüzü", "Bildirimler & sohbet"], True)
    box(d, 410, 100, 280, 120, ["API & Gerçek Zamanlı", "Express + Socket.IO", "REST /api", "Anlık olay yayınları"])
    box(d, 760, 100, 280, 120, ["Veri Katmanı", "Prisma ORM", "SQLite / PostgreSQL", "Dosya depolama"])
    arrow(d, 340, 160, 410, 160)
    arrow(d, 690, 160, 760, 160)
    box(d, 230, 300, 640, 100, ["Ortak Paket (Shared)", "Rol × İzin matrisi · Doğrulama şemaları · Socket olay sözleşmeleri", "Tek kaynak: web ve sunucu aynı kuralları kullanır"], True)
    arrow(d, 550, 220, 550, 300)
    box(d, 80, 460, 220, 90, ["Kimlik", "JWT çerezleri", "CSRF koruması"])
    box(d, 340, 460, 220, 90, ["Yetki", "Platform RBAC", "Sunucu rolleri"])
    box(d, 600, 460, 220, 90, ["Odalar", "Kanal / Sunucu", "Takım / Kullanıcı"])
    box(d, 860, 460, 180, 90, ["Entegrasyon", "CTFd API", "Webhook"])
    return save_img(img, "flow_architecture.png")


def chart_auth() -> Path:
    w, h = 1100, 520
    img = Image.new("RGB", (w, h), BG)
    d = ImageDraw.Draw(img)
    title_bar(d, w, "Kimlik Doğrulama Akışı", "Güvenli oturum açma ve gerçek zamanlı bağlantı")
    steps = [
        (40, 120, ["1. CSRF", "Token al"]),
        (240, 120, ["2. Giriş", "Kullanıcı / e-posta"]),
        (440, 120, ["3. Çerezler", "Access + Refresh"]),
        (640, 120, ["4. Socket", "Çerezle bağlan"]),
        (840, 120, ["5. Hazır", "Presence & sohbet"]),
    ]
    for x, y, lines in steps:
        box(d, x, y, 170, 90, lines, accent=True)
    for i in range(4):
        arrow(d, 210 + i * 200, 165, 240 + i * 200, 165)
    box(d, 240, 280, 620, 140, [
        "Oturum yenileme",
        "Access token süresi dolunca API 401 döner",
        "İstemci otomatik olarak /auth/refresh çağırır",
        "Yeni çerezlerle istek tekrarlanır — kullanıcı fark etmez",
        "Çıkış / tüm cihazlardan çıkış desteklenir",
    ])
    return save_img(img, "flow_auth.png")


def chart_competitor() -> Path:
    w, h = 1100, 700
    img = Image.new("RGB", (w, h), BG)
    d = ImageDraw.Draw(img)
    title_bar(d, w, "Yarışmacı Yolculuğu", "Katılımdan sohbet, destek ve skor tablosuna")
    box(d, 420, 90, 260, 70, ["Kayıt / Giriş"], True)
    arrow(d, 550, 160, 550, 190)
    box(d, 380, 190, 340, 80, ["Sunucu üyeliği?", "Davet kodu ile katıl veya mevcut sunucu"])
    arrow(d, 550, 270, 550, 300)
    box(d, 400, 300, 300, 70, ["Kanal sohbeti", "Takım / genel kanallar"], True)
    # branches
    arrow(d, 400, 335, 180, 400)
    arrow(d, 700, 335, 920, 400)
    arrow(d, 550, 370, 550, 400)
    box(d, 40, 400, 280, 100, ["Destek talebi", "Kategori seç · Açıklama yaz", "Dosya ekle · Gönder"])
    box(d, 410, 400, 280, 100, ["Takım alanı", "Özel kanal / kategori", "Takım arkadaşlarıyla koordine ol"])
    box(d, 780, 400, 280, 100, ["Skor tablosu", "CTFd canlı sıralama", "Kategori grafikleri"])
    arrow(d, 180, 500, 180, 540)
    box(d, 40, 540, 280, 100, ["Bildirimler", "Yanıt / durum değişimi", "Bahsetme & duyurular"])
    return save_img(img, "flow_competitor.png")


def chart_tickets() -> Path:
    w, h = 1200, 720
    img = Image.new("RGB", (w, h), BG)
    d = ImageDraw.Draw(img)
    title_bar(d, w, "Destek Talebi Yaşam Döngüsü", "Açılıştan çözüm ve yeniden açılışa")
    statuses = [
        (40, 100, ["AÇIK", "Yeni talep"]),
        (260, 100, ["İŞLEMDE", "Üstlenildi"]),
        (480, 100, ["BEKLİYOR", "Yarışmacı yanıtı"]),
        (700, 100, ["ÇÖZÜLDÜ", "Sonuçlandı"]),
        (920, 100, ["KAPATILDI", "Arşiv"]),
    ]
    for x, y, lines in statuses:
        box(d, x, y, 190, 80, lines, accent=True)
    for i in range(4):
        arrow(d, 230 + i * 220, 140, 260 + i * 220, 140)
    # reopen loop
    d.line((1015, 180, 1015, 250, 135, 250, 135, 180), fill=ARROW, width=3)
    arrow(d, 135, 200, 135, 180)
    d.text((420, 230), "Yeniden açıldı → tekrar Açık / İşlemde", font=font(13), fill=MUTED)

    box(d, 40, 300, 350, 160, [
        "Yarışmacı",
        "• Talep oluşturur",
        "• Genel yorum yazar",
        "• Ek dosya yükler",
        "• Çözülen talebi yeniden açar",
        "• Kendi / takım taleplerini görür",
    ])
    box(d, 420, 300, 360, 160, [
        "Organizatör (Mod+)",
        "• Tüm talepleri görür",
        "• Üstlenir / atar",
        "• İç not yazar (gizli)",
        "• Durum / öncelik / etiket",
        "• İlişki & duyuru oluşturur",
    ], True)
    box(d, 810, 300, 350, 160, [
        "Talep Operasyonu",
        "• Son 5/15/30 dk gelişler",
        "• Atanmamış kuyruk",
        "• Takım / görev yükü",
        "• Çoklu takım olayları",
        "• Ortalama yanıt raporları",
    ])
    box(d, 200, 510, 800, 140, [
        "Bildirim & gerçek zamanlı yayın",
        "Yeni talep → sunucu personel odası · Yanıt → muhatap kullanıcı odası",
        "Durum / atama / kritik öncelik → ilgili taraflara anlık bildirim",
        "Takım odası yalnızca aynı yarışma sunucusu kapsamında kullanılır (sızıntı önleme)",
    ])
    return save_img(img, "flow_tickets.png")


def chart_admin() -> Path:
    w, h = 1100, 680
    img = Image.new("RGB", (w, h), BG)
    d = ImageDraw.Draw(img)
    title_bar(d, w, "Yarışma Kurulum Akışı (Yönetici)", "Sıfırdan canlı operasyona")
    steps = [
        (80, 100, ["1", "Sunucu oluştur", "İsim · ikon · tema"]),
        (400, 100, ["2", "Roller & izinler", "Varsayılan üye rolü"]),
        (720, 100, ["3", "Kategori/Kanal", "Staff · genel · takım"]),
        (80, 280, ["4", "CTFd içe aktar", "Takım + görev CSV"]),
        (400, 280, ["5", "CTFd bağlantı", "URL + API token"]),
        (720, 280, ["6", "Talep ayarları", "Önek · politika"]),
        (240, 460, ["7", "Davetleri dağıt", "Kod / süre / limit"]),
        (560, 460, ["8", "Canlı operasyon", "Skor · talep · duyuru"]),
    ]
    for x, y, lines in steps:
        box(d, x, y, 280, 100, [f"{lines[0]}. {lines[1]}", lines[2]], accent=True)
    arrow(d, 360, 150, 400, 150)
    arrow(d, 680, 150, 720, 150)
    arrow(d, 860, 200, 860, 280)
    arrow(d, 720, 330, 680, 330)
    arrow(d, 400, 330, 360, 330)
    arrow(d, 220, 380, 220, 460)
    arrow(d, 520, 510, 560, 510)
    return save_img(img, "flow_admin.png")


def chart_messaging() -> Path:
    w, h = 1100, 560
    img = Image.new("RGB", (w, h), BG)
    d = ImageDraw.Draw(img)
    title_bar(d, w, "Mesajlaşma & Gerçek Zamanlı Yayın", "Yazma göstergesinden ekrana")
    box(d, 40, 120, 240, 120, ["Kullanıcı A", "Mesaj yazar", "Enter ile gönderir"], True)
    arrow(d, 280, 180, 340, 180)
    box(d, 340, 120, 260, 120, ["API", "Doğrula · kaydet", "Ekleri sakla"])
    arrow(d, 600, 180, 660, 180)
    box(d, 660, 120, 380, 120, ["Socket.IO yayını", "ch:{kanalId} odasına", "message:created olayı"])
    box(d, 160, 320, 780, 160, [
        "Alıcı tarafta",
        "SocketBridge olayı yakalar → TanStack Query önbelleği güncellenir",
        "Liste anında kayar · bahsetme varsa toast + isteğe bağlı tarayıcı bildirimi",
        "Aynı anda typing:start/stop ve presence (çevrimiçi / meşgul / DND) akar",
    ])
    arrow(d, 850, 240, 850, 320)
    return save_img(img, "flow_messaging.png")


def chart_roles() -> Path:
    w, h = 1100, 520
    img = Image.new("RGB", (w, h), BG)
    d = ImageDraw.Draw(img)
    title_bar(d, w, "Rol Hiyerarşisi (Platform)", "Aşağıdan yukarıya artan yetki")
    roles = [
        ("Gözlemci", "İzleme"),
        ("Seyirci", "Kayıtlı izleyici"),
        ("Operatör", "Takım üyesi"),
        ("Kaptan", "Takım yönetimi"),
        ("Moderatör", "Personel"),
        ("Görev Yazarı", "Challenge"),
        ("Yarışma Yön.", "Organizasyon"),
        ("Süper Admin", "Tam yetki"),
    ]
    x = 30
    for i, (name, sub) in enumerate(roles):
        box(d, x, 140 if i % 2 == 0 else 260, 120, 90, [name, sub], accent=i >= 4)
        if i < len(roles) - 1:
            arrow(d, x + 120, 185 if i % 2 == 0 else 305, x + 140, 245 if i % 2 == 0 else 185)
        x += 135
    d.text(
        (40, 400),
        "Ayrıca her yarışma sunucusunda Discord tarzı sunucu rolleri (kanal yönet, yasakla, davet et…) vardır.",
        font=font(14),
        fill=MUTED,
    )
    d.text(
        (40, 430),
        "Personel eşiği: Moderatör ve üzeri · Yönetici eşiği: Yarışma Yöneticisi ve üzeri.",
        font=font(14),
        fill=MUTED,
    )
    return save_img(img, "flow_roles.png")


# ---------- Word helpers ----------

def set_run_font(run, size=11, bold=False, color=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading_styled(doc: Document, text: str, level: int):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E) if level <= 1 else RGBColor(0x13, 0x4E, 0x4A)
    return p


def body(doc: Document, text: str, *, bold_first=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if bold_first and ":" in text:
        label, rest = text.split(":", 1)
        r1 = p.add_run(label + ":")
        set_run_font(r1, 11, True)
        r2 = p.add_run(rest)
        set_run_font(r2, 11)
    else:
        r = p.add_run(text)
        set_run_font(r, 11)
    return p


def bullets(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r, 11)
        p.paragraph_format.space_after = Pt(2)


def add_image(doc: Document, path: Path, width_inches: float = 6.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_inches))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cap.add_run(f"Şekil: {path.stem.replace('_', ' ')}")
    set_run_font(cr, 9, color=RGBColor(0x64, 0x74, 0x8B))


def table_kv(doc: Document, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for run in p.runs:
                set_run_font(run, 10, True, RGBColor(0x0F, 0x17, 0x2A))
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
            for p in cells[ci].paragraphs:
                for run in p.runs:
                    set_run_font(run, 10)
    doc.add_paragraph()


def build_document(charts: dict[str, Path]) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    # Cover
    for _ in range(3):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("SYDT Chat")
    set_run_font(r, 36, True, RGBColor(0x0F, 0x76, 0x6E))

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run("CTF Operasyon Konsolu")
    set_run_font(r, 20, True, RGBColor(0x33, 0x41, 0x55))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Ürün Tanıtım ve Kullanıcı Kılavuzu")
    set_run_font(r, 14, color=RGBColor(0x64, 0x74, 0x8B))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        "Sürüm 0.1 · Gerçek zamanlı iletişim · Destek talepleri · CTFd skor tablosu\n"
        "Bu belge, sistemi son kullanıcıya bir ürün olarak tanıtır; tüm ana özellikleri ve işleyiş akışlarını kapsar."
    )
    set_run_font(r, 11, color=RGBColor(0x47, 0x55, 0x69))

    doc.add_page_break()

    # TOC-like outline
    add_heading_styled(doc, "İçindekiler", 1)
    toc = [
        "1. Ürüne Giriş — SYDT Chat Nedir?",
        "2. Kime Hitap Eder?",
        "3. Sistemin Genel Mimarisi",
        "4. İlk Adımlar: Hesap, Giriş ve Güvenlik",
        "5. Arayüz Turu",
        "6. Yarışma Sunucuları",
        "7. Kategoriler ve Kanallar",
        "8. Mesajlaşma ve Dosyalar",
        "9. Özel Mesajlar ve Arkadaşlar",
        "10. Takımlar",
        "11. Davetler, Roller ve Moderasyon",
        "12. Destek Talepleri (Tickets)",
        "13. Talep Operasyon Paneli",
        "14. Görevler (Challenges)",
        "15. Skor Tablosu ve CTFd",
        "16. Olaylar (Incidents)",
        "17. Duyurular ve Bildirimler",
        "18. Anketler ve Platform Sayfaları",
        "19. Yönetim Paneli",
        "20. Roller, İzinler ve Güvenlik",
        "21. Klavye Kısayolları ve İpuçları",
        "22. Uçtan Uca Senaryolar (Akış Şemaları)",
        "23. Bilerek Sonraki Aşamaya Bırakılanlar",
        "24. Sözlük",
    ]
    for item in toc:
        body(doc, item)

    doc.add_page_break()

    # 1
    add_heading_styled(doc, "1. Ürüne Giriş — SYDT Chat Nedir?", 1)
    body(
        doc,
        "SYDT Chat, siber güvenlik CTF (Capture The Flag) yarışmalarının operasyonel omurgası olmak üzere tasarlanmış "
        "gerçek zamanlı bir iletişim ve operasyon platformudur. Tek cümleyle: Discord benzeri sunucu/kanal deneyimini, "
        "yarışma destek masası (ticket), canlı skor tablosu, olay yönetimi ve CTFd entegrasyonu ile birleştirir.",
    )
    body(
        doc,
        "Ürün adı ve marka varsayılanları platformda şöyle konumlanır: platform başlığı “SYDT”, slogan “CTF Operations Console”, "
        "çalışma alanı başlığı “SYDT Chat”. Arayüz varsayılan dili Türkçedir; İngilizce sözlük de mevcuttur.",
    )
    body(doc, "SYDT Chat’in vaat ettiği deneyim:", bold_first=False)
    bullets(
        doc,
        [
            "Yarışma sırasında takımlar ve organizatörler aynı “çatı” altında koordine olur.",
            "Sohbet, bahsetme, tepki, dosya ekleri ve yazıyor göstergeleri anlık akar.",
            "Teknik sorunlar destek talebi olarak takip edilir; personel üstlenir, iç not yazar, çözer.",
            "CTFd skorları canlı çekilir; sıralama ve kategori dağılımı ekranda canlanır.",
            "Duyurular, anketler, yasaklar, davetler ve denetim kayıtları operasyonu güvende tutar.",
        ],
    )

    # 2
    add_heading_styled(doc, "2. Kime Hitap Eder?", 1)
    table_kv(
        doc,
        ["Kullanıcı tipi", "Ne için kullanır?"],
        [
            ["Yarışmacı / Operatör", "Takım sohbeti, görev tartışması, destek talebi, skor takibi"],
            ["Takım kaptanı", "Takım yönetimi, üye koordinasyonu, özel kanallar"],
            ["Moderatör / Personel", "Staff kanalları, talep kuyruğu, susturma, olay izleme"],
            ["Görev yazarı", "Challenge içeriği ve tartışma kanalları"],
            ["Yarışma yöneticisi", "Sunucu kurulumu, duyurular, CTFd, talep politikası"],
            ["Süper yönetici", "Platform markası, kullanıcı/moderasyon, anket, denetim"],
            ["Gözlemci / Seyirci", "Sınırlı izleme; seyirci kendi taleplerini açabilir"],
        ],
    )

    # 3
    add_heading_styled(doc, "3. Sistemin Genel Mimarisi", 1)
    body(
        doc,
        "Son kullanıcı bir web tarayıcısından çalışır. Arka planda üç ana katman vardır: Web arayüzü, API + gerçek zamanlı katman "
        "ve veri katmanı. Ortak bir paket, rol-izin kurallarını ve soket olay sözleşmesini tek kaynakta tutar; böylece web ve sunucu "
        "aynı dili konuşur.",
    )
    add_image(doc, charts["architecture"])
    bullets(
        doc,
        [
            "Web: Next.js 15, React 19, TanStack Query (sunucu verisi), Zustand (hafif istemci durumu), Socket.IO istemcisi.",
            "Sunucu: Express REST API, Socket.IO odaları, Prisma ile veritabanı, Argon2 parola özeti, JWT çerez oturumu.",
            "Veri: Geliştirmede SQLite; şema PostgreSQL’e taşınabilir. Yüklenen dosyalar ayrı bir depolama dizininde tutulur.",
            "Sağlık: API /health uç noktası ile ayakta olup olmadığı kontrol edilir.",
        ],
    )

    # 4
    add_heading_styled(doc, "4. İlk Adımlar: Hesap, Giriş ve Güvenlik", 1)
    add_heading_styled(doc, "4.1 Kayıt ve giriş", 2)
    body(
        doc,
        "Kullanıcı /login veya /register sayfalarından ilerler. Girişte kullanıcı adı veya e-posta kullanılabilir. "
        "İlk istekten önce sistem CSRF (çapraz site istek sahteciliği) koruması için bir token alır; ardından giriş bilgileri gönderilir.",
    )
    add_heading_styled(doc, "4.2 Oturum modeli", 2)
    bullets(
        doc,
        [
            "Kısa ömürlü erişim (access) JWT’si httpOnly çerezde tutulur.",
            "Uzun ömürlü yenileme (refresh) JWT’si ayrı bir httpOnly çerezdedir.",
            "Access süresi dolunca istemci sessizce yeniler; kullanıcıyı tekrar şifre sormaya zorlamaz.",
            "Socket.IO el sıkışması da aynı çerezi kullanır — sohbet ve API aynı kimlikle yürür.",
            "Tek oturumdan veya tüm cihazlardan çıkış desteklenir.",
        ],
    )
    add_image(doc, charts["auth"])
    add_heading_styled(doc, "4.3 Zorunlu parola değişimi", 2)
    body(
        doc,
        "Bir sunucu “katıldığında parola değiştir” politikası açıksa, PasswordChangeGate kullanıcıyı diğer ekranlara "
        "geçmeden önce yeni parola belirlemeye yönlendirir. Bu, paylaşılan / içe aktarılan hesaplarda güvenlik katmanıdır.",
    )
    add_heading_styled(doc, "4.4 Güvenlik özeti (kullanıcıya yansıyan)", 2)
    bullets(
        doc,
        [
            "Parolalar Argon2 ile saklanır; düz metin tutulmaz.",
            "Mutasyon isteklerinde X-CSRF-Token çift gönderim koruması vardır.",
            "Hız sınırları (rate limit) giriş, talep oluşturma, yanıt ve dosya indirme gibi uçlarda kötüye kullanımı keser.",
            "Yetkisiz destek talebi erişiminde sistem bilerek 404 döner — tahmin ederek talep numarası avlamayı zorlaştırır.",
        ],
    )

    # 5
    add_heading_styled(doc, "5. Arayüz Turu", 1)
    body(
        doc,
        "Uygulama kabuğu (app layout) soldan sağa şu omurgayı kurar: ServerRail (sunucu şeridi) → ChannelSidebar "
        "(kanal / DM listesi) → ana içerik alanı. Üst çubukta arama/komut paleti, bildirimler ve profil kısayolları yer alır.",
    )
    table_kv(
        doc,
        ["Bileşen", "Ne işe yarar?"],
        [
            ["ServerRail", "Üye olunan yarışma sunucuları arasında geçiş; oluştur / ayarlar"],
            ["ChannelSidebar", "Kategoriler, kanallar, DM konuşmaları, bildirim rozetleri"],
            ["TopBar", "Bağlam başlığı, bildirim zili, kullanıcı menüsü"],
            ["Command Palette (Ctrl/⌘+K)", "Hızlı gezinme ve komut arama"],
            ["SurveyGate", "Zorunlu anket varken devamı engeller"],
            ["PasswordChangeGate", "Zorunlu parola değişimini uygular"],
            ["SocketBridge", "Anlık olayları arayüz önbelleğine yansıtır"],
        ],
    )
    body(doc, "Ana rotalar (ekranlar):", bold_first=False)
    bullets(
        doc,
        [
            "/ — tercih edilen kanala yönlendirme veya sunucusuz karşılama",
            "/channels/[id] — kanal sohbeti",
            "/tickets, /tickets/[id], /tickets/ops — destek ve operasyon",
            "/scoreboard — skor tablosu",
            "/challenges, /challenges/[id] — görevler",
            "/incidents — olaylar",
            "/team — takımım",
            "/friends — arkadaşlar",
            "/notifications — bildirim merkezi",
            "/settings/profile ve /settings/notifications — ayarlar",
            "/admin — yönetim paneli",
            "/invite/[code] — davet kabul",
            "/p/[slug] — platform sayfaları · /legal/* — kullanım koşulları",
        ],
    )

    # 6
    add_heading_styled(doc, "6. Yarışma Sunucuları", 1)
    body(
        doc,
        "SYDT Chat’te her CTF operasyonu bir “sunucu” (ChatServer) olarak yaşar. Kullanıcı birden fazla sunucuya üye olabilir; "
        "şerit üzerinden aktif sunucuyu değiştirir. Sunucu; ikon, banner, tema renkleri, katılımda parola değiştirme zorunluluğu "
        "ve e-posta bildirim bayrakları gibi kimlik/tercih alanları taşır.",
    )
    add_heading_styled(doc, "6.1 Sunucu ayarları sekmeleri", 2)
    bullets(
        doc,
        [
            "Genel Bakış — kimlik, görünüm, temel politikalar",
            "Kategoriler — ağaç yapısı, personel-only, takıma özel çalışma alanları",
            "Kanallar — tür, kilit, logo, bildirim tercihi",
            "CSV aktarımı — CTFd / toplu veri",
            "Roller — renkli roller ve JSON izinler",
            "Emojiler — özel emoji yönetimi",
            "Üyeler — üyelik, takma ad, rol atama",
            "Takımlar — sunucu içi takım listesi",
            "Davetler — kod, kullanım limiti, son kullanma",
            "Yasaklananlar — ban / unban",
            "Destek talepleri — kategori, etiket, yarışma ayarları",
            "Entegrasyonlar — CTFd URL/token, webhook’lar",
        ],
    )
    add_heading_styled(doc, "6.2 Sunucusuz kullanıcı", 2)
    body(
        doc,
        "Henüz hiçbir sunucuda olmayan kullanıcı NoServerOnboarding deneyimini görür: davet kodu ile katılma veya "
        "yetkisi varsa sunucu oluşturma yolu sunulur.",
    )

    # 7
    add_heading_styled(doc, "7. Kategoriler ve Kanallar", 1)
    body(
        doc,
        "Kanallar kategoriler altında gruplanır. Bazı kategoriler yalnızca personel içindir; bazıları belirli bir takıma "
        "bağlı özel çalışma alanıdır (Category.teamId). Kanallar kilitlenebilir; kanal veya sunucu bazında bildirim modu seçilebilir.",
    )
    add_heading_styled(doc, "7.1 Kanal türleri", 2)
    table_kv(
        doc,
        ["Tür", "Anlamı"],
        [
            ["GLOBAL_ANNOUNCEMENT", "Genel duyuru kanalı"],
            ["GLOBAL_CHAT", "Herkese açık sohbet"],
            ["TEAM_SUPPORT", "Takım destek / yardım"],
            ["CHALLENGE_DISCUSSION", "Görev tartışması"],
            ["INCIDENT", "Olay koordinasyonu"],
            ["VOICE_PLACEHOLDER", "Sesli kanal yer tutucu (henüz medya yok)"],
            ["STAFF_COORD", "Personel koordinasyonu"],
            ["ORGANIZER_INTERNAL", "Organizatör iç kanalı"],
            ["LOGS", "Operasyon günlükleri"],
            ["TEAM_PRIVATE", "Takıma özel özel kanal"],
            ["DIRECT_MESSAGE", "Özel mesaj (DM) kanalı"],
        ],
    )

    # 8
    add_heading_styled(doc, "8. Mesajlaşma ve Dosyalar", 1)
    body(
        doc,
        "Kanal sohbeti sanal kaydırmalı (virtuoso) bir listedir: yukarı kaydırınca eski mesajlar sayfa sayfa gelir. "
        "Markdown, kod blokları (Shiki), IOC/IP/hash/URL/JSON/terminal vurguları desteklenir. Mesajlara yanıt verilir, "
        "düzenlenir, silinir; tepki (reaction) eklenir; sabitlenen mesajlar panelde toplanır.",
    )
    bullets(
        doc,
        [
            "Gönder: Enter · Yeni satır: Shift+Enter",
            "Dosya: besteci üzerindeki ataş veya sürükle-bırak (multipart)",
            "Bahsetme (@): otomatik tamamlama; hedefe bildirim düşer",
            "Yazıyor göstergesi ve çevrimiçi durumu (presence) anlıktır",
            "Dosya indirme: /api/files/:id — çerez kimliği + kanal/talep erişim kontrolü",
        ],
    )
    add_image(doc, charts["messaging"])

    # 9
    add_heading_styled(doc, "9. Özel Mesajlar ve Arkadaşlar", 1)
    body(
        doc,
        "Arkadaşlık isteği gönderilir, kabul veya reddedilir. DM açmak için POST /api/direct/open kullanılır. "
        "Kritik kural: personel herkese DM atabilir; yarışmacılar yalnızca ortak yarışma takımında oldukları kişilerle "
        "DM kurabilir (ServerTeamMember kesişimi). Bu, yarışmalar arası istenmeyen iletişimi keser.",
    )

    # 10
    add_heading_styled(doc, "10. Takımlar", 1)
    body(
        doc,
        "Takımlar sunucu kapsamındadır. Kaptan, davet kodu, logo, kadro ve CTFd takım kimliği bağlantısı yönetilir. "
        "Takıma özel kategori/kanal açılabilir. /team ekranı “Takımım” görünümünü sunar; yöneticiler admin ve sunucu "
        "ayarlarından takımları yönetir.",
    )

    # 11
    add_heading_styled(doc, "11. Davetler, Roller ve Moderasyon", 1)
    add_heading_styled(doc, "11.1 Davetler", 2)
    bullets(
        doc,
        [
            "Davet kodu: maksimum kullanım, son kullanma tarihi, isteğe bağlı iniş kanalı",
            "Kabul sayfası: /invite/[code] — “Sunucuya katıl”",
            "Yasaklı kullanıcı davetle geri dönemez",
        ],
    )
    add_heading_styled(doc, "11.2 Sunucu rolleri", 2)
    body(
        doc,
        "Her sunucuda renkli, isimli roller vardır. İzinler JSON olarak tutulur: manageChannels, manageServer, "
        "manageRoles, banMembers, kickMembers, inviteOthers, sendMessage, uploadFiles, viewChannel vb. Üyeler birden "
        "fazla role sahip olabilir; izinler birleşir. Platform ADMIN_DASHBOARD yetkisi sunucu izin kontrollerini aşabilir.",
    )
    add_heading_styled(doc, "11.3 Yasaklar ve raporlar", 2)
    bullets(
        doc,
        [
            "Sunucu banı: sebep ile yasaklama / kaldırma",
            "Platform susturma / yasaklama: yönetici panelinden",
            "Kullanıcı raporları: /api/reports → admin incelemesi",
        ],
    )

    # 12
    add_heading_styled(doc, "12. Destek Talepleri (Tickets)", 1)
    body(
        doc,
        "Destek talepleri SYDT’nin kalbidir. Her talep bir yarışma sunucusuna bağlıdır (serverId). Numaralandırma "
        "yarışma ayarlarındaki önek ile üretilir (ör. LS26-000001). Talep politikası OPEN (herkes), ORGANIZER_ONLY "
        "veya CLOSED olabilir; yarışma durumu UPCOMING / LIVE / PAUSED / FINISHED ile uyumlu yönetilir.",
    )
    add_heading_styled(doc, "12.1 Durumlar", 2)
    table_kv(
        doc,
        ["Kod", "Türkçe", "Anlam"],
        [
            ["OPEN", "Açık", "Yeni veya yeniden açılmış, bekleyen iş"],
            ["IN_PROGRESS", "İşlemde", "Personel üzerinde çalışıyor"],
            ["WAITING_FOR_COMPETITOR", "Bekliyor", "Yarışmacıdan bilgi bekleniyor"],
            ["RESOLVED", "Çözüldü", "Sonuçlandı"],
            ["CLOSED", "Kapatıldı", "Arşivlendi"],
            ["REOPENED", "Yeniden açıldı", "Çözüm sonrası tekrar gündemde"],
        ],
    )
    add_heading_styled(doc, "12.2 Öncelik / önem", 2)
    body(doc, "Düşük · Orta · Yüksek · Kritik. Sıralama için sayısal priorityRank kullanılır.")
    add_heading_styled(doc, "12.3 Varsayılan kategoriler", 2)
    bullets(
        doc,
        [
            "Görev sorunu",
            "Görev açıklaması",
            "Teknik sorun",
            "Altyapı",
            "Bayrak gönderim sorunu",
            "Skor sorunu",
            "Hesap / Takım",
            "Yarışma",
            "Diğer",
        ],
    )
    body(doc, "Bazı kategoriler bir challenge (görev) seçimini zorunlu kılar.")
    add_heading_styled(doc, "12.4 Kim ne yapabilir?", 2)
    table_kv(
        doc,
        ["Aktör", "Yetkiler"],
        [
            ["Yarışmacı", "Politika izin veriyorsa oluşturur; yorum; ek; yeniden açma; kendi/takım talepleri"],
            ["Moderatör+", "Tümünü görür; üstlenir; atar; durum/öncelik/etiket; iç not; toplu işlem; ilişki"],
            ["Yarışma yöneticisi+", "Kategori/etiket/yarışma meta; duyuru üretme"],
        ],
    )
    add_heading_styled(doc, "12.5 Yaşam döngüsü adımları", 2)
    bullets(
        doc,
        [
            "Oluştur: TicketComposer → sunucu ticket API (multipart ekler) → kod üretimi → personele bildirim",
            "Yanıt: genel yorum veya INTERNAL_NOTE (iç not — yalnızca personel)",
            "Üstlen (claim): kendini ata, genelde İşlemde’ye geç",
            "Güncelle: durum, öncelik, atanan, kategori, etiket, ilişkili görev",
            "Yeniden aç: gerekçeli reopen",
            "İlişkiler: RELATED / DUPLICATE / PARENT / CHILD (ters kenar da tutulur)",
            "Duyuru: ticket’tan announce — gövde talepten otomatik kopyalanmaz; bilinçli yazılır",
            "Okundu imleci: unread sayaçları için read endpoint",
            "Toplu işlem: uygun izinle çoklu talep üzerinde aksiyon",
        ],
    )
    body(
        doc,
        "Ekler: talep olayına en fazla 5 dosya; ekran görüntüsü, log, PDF, pcap gibi operasyonel türler desteklenir "
        "(arayüzde Türkçe ipucu ile).",
    )
    add_image(doc, charts["tickets"])

    # 13
    add_heading_styled(doc, "13. Talep Operasyon Paneli", 1)
    body(
        doc,
        "/tickets/ops ekranı “Talep operasyonu” adıyla personelin nabzını tutar. Son 5 / 15 / 30 dakikalık gelişler, "
        "atanmamış kuyruk, takım / görev / kategori / atanan bazında yük, çoklu takım olayları ve ortalama ilk yanıt / "
        "çözüm süreleri raporlanır. Bu panel, büyük CTF’lerde destek masasını yönetilebilir kılar.",
    )

    # 14
    add_heading_styled(doc, "14. Görevler (Challenges)", 1)
    body(
        doc,
        "/challenges listesi ve /challenges/[id] detayı ile görevler filtrelenir, ipuçları ve çözümler incelenir, "
        "bayrak gönderilir, first blood kutlanır. Görev tartışma kanalına bağlantı kurulabilir. CTFd challenges.csv "
        "içe aktarımı hem challenge kayıtlarını hem ilgili kanalları üretebilir.",
    )

    # 15
    add_heading_styled(doc, "15. Skor Tablosu ve CTFd", 1)
    body(
        doc,
        "Sunucuya CTFd taban URL’si ve erişim belirteci girildiğinde skor tablosu CTFd’den canlı çekilir. Yerel takımlar "
        "ctfdTeamId veya isim eşlemesiyle eşlenir. Kategori grafikleri ve üst görev istatistikleri admin API istatistikleriyle "
        "beslenir. CTFd yoksa yerel skor tablosu / zaman çizelgesi hâlâ mevcuttur. Eski yerel skor anlık görüntüsü periyodik "
        "yayınlanabilir; arayüz sıralama değişimlerini hareketle gösterir.",
    )
    bullets(
        doc,
        [
            "Ekran: /scoreboard — “CTFd skor tablosu” / “Canlı skor tablosu”",
            "Bileşenler: CtfdRankRow, CtfdCategoryChart",
            "Entegrasyon paneli: bağlantı testi",
            "CSV: takım ve görev içe aktarımı",
            "Giden URL’ler güvenlik süzgecinden geçer (url-safety)",
        ],
    )

    # 16
    add_heading_styled(doc, "16. Olaylar (Incidents)", 1)
    body(
        doc,
        "Olaylar (Incident) yarışma çapında veya platform çapında kritik durumları takip eder. Durum akışı: "
        "OPEN → INVESTIGATING → MITIGATED → RESOLVED → POSTMORTEM. Önem: INFO … CRITICAL. Tercihen serverId ile "
        "kapsamlanır; eski kayıtlarda staff odasına yayın yapılabilir. Liste ve detay sayfaları /incidents altındadır.",
    )

    # 17
    add_heading_styled(doc, "17. Duyurular ve Bildirimler", 1)
    add_heading_styled(doc, "17.1 Duyurular", 2)
    body(
        doc,
        "Sunucu kapsamlı duyurular zamanlanabilir, yayınlanabilir, süresi dolabilir. Ticket kaynaklı duyuru üretilebilir. "
        "Entegrasyon webhook’ları announcement.publish olayını dış sistemlere iletebilir. Anlık olay: announcement:new.",
    )
    add_heading_styled(doc, "17.2 Bildirimler", 2)
    bullets(
        doc,
        [
            "Türler: talep oluşturuldu / yanıt / durum / atandı / çözüldü / yeniden açıldı / kritik / çoklu takım; duyuru; bahsetme; sistem",
            "Kalıcı gelen kutusu veritabanında; anlık katman Zustand + SocketBridge",
            "Ses efektleri (WebAudio), sonner toast, tarayıcı Notification API",
            "DND ve ses anahtarları; kanal/sunucu notifyMode tercihleri",
            "Ekran: /notifications ve üst çubuk popover",
        ],
    )

    # 18
    add_heading_styled(doc, "18. Anketler ve Platform Sayfaları", 1)
    body(
        doc,
        "Yöneticiler tek seçimli / çok seçimli / serbest metin anketleri oluşturur. “Devam etmeden önce zorunlu” "
        "işaretli anketler SurveyGate ile tüm uygulamayı bloke eder. PlatformPage kayıtları /p/[slug] altında CMS benzeri "
        "içerik sunar. Yasal metinler /legal altında EN/TR kabul edilebilir kullanım sayfaları olarak bulunur.",
    )

    # 19
    add_heading_styled(doc, "19. Yönetim Paneli", 1)
    body(doc, "/admin paneli sekmeleri (Türkçe arayüz etiketleriyle):")
    bullets(
        doc,
        [
            "Genel bakış — KPI kartları, aktivite akışı, çevrimiçi personel",
            "Takımlar — platform/takım yönetimi",
            "Sistem ve ayarlar — yükleme limitleri, kayıt varsayılan sunucusu vb.",
            "Platform markası — başlık, slogan, görseller",
            "Duyurular ve toplu DM — AdminBroadcast / kitlesel iletişim",
            "Anketler",
            "Sayfalar — PlatformPage",
            "Moderasyon — raporlar, kullanıcı işlemleri",
            "İşlem kayıtları — AuditLog",
        ],
    )
    add_image(doc, charts["admin"])

    # 20
    add_heading_styled(doc, "20. Roller, İzinler ve Güvenlik", 1)
    add_heading_styled(doc, "20.1 Platform rolleri", 2)
    add_image(doc, charts["roles"])
    body(
        doc,
        "Personel eşiği Moderatör ve üzeridir (isStaff). Yönetici eşiği Yarışma Yöneticisi ve üzeridir (isAdmin). "
        "Bir kullanıcı yarışma takımına katıldığında platform personel rolü platformRole alanında korunur; etkili izin "
        "rolü buna göre hesaplanır.",
    )
    add_heading_styled(doc, "20.2 Çift katmanlı yetki", 2)
    bullets(
        doc,
        [
            "Platform: Role × Permission matrisi (packages/shared) — sunucu requirePermission + soket korumaları + web <Can>/useCan",
            "Sunucu: Discord tarzı permissionsJson — üyelik ve kanal eylemleri",
            "Ticket izolasyonu: ServerTeamMember; yetkisiz erişimde 404",
            "İç notlar: yalnızca personel soketine ve filtrelenmiş sorguya",
            "Denetim kaydı: hassas eylemler AuditLog’a yazılır",
        ],
    )
    add_heading_styled(doc, "20.3 Gerçek zamanlı odalar", 2)
    table_kv(
        doc,
        ["Oda", "Amaç"],
        [
            ["global", "Presence, genel skor"],
            ["user:{id}", "Bildirim, bahsetme, DM yenileme"],
            ["ch:{id}", "Kanal mesajları / typing"],
            ["srv:{id}", "Üyelik olayları / duyurular"],
            ["srv:{id}:staff", "Talep/olay personel yayını"],
            ["srv:{id}:team:{teamId}", "Yarışma kapsamlı takım talebi"],
            ["staff", "Eski, sunucusuz olaylar"],
        ],
    )

    # 21
    add_heading_styled(doc, "21. Klavye Kısayolları ve İpuçları", 1)
    table_kv(
        doc,
        ["Eylem", "Kısayol"],
        [
            ["Komut paletini aç", "Ctrl / ⌘ + K"],
            ["Palet / modal kapat", "Esc"],
            ["Mesaj gönder", "Enter"],
            ["Sohbette yeni satır", "Shift + Enter"],
            ["Palette ara", "Yazmaya başla"],
        ],
    )

    # 22
    add_heading_styled(doc, "22. Uçtan Uca Senaryolar (Akış Şemaları)", 1)
    add_heading_styled(doc, "22.1 Yarışmacı günü", 2)
    body(
        doc,
        "Sabah kayıt olur veya giriş yapar → davetle sunucuya katılır → genel ve takım kanallarında koordine olur → "
        "bir görevde takılırsa destek talebi açar → yanıtları bildirimden takip eder → arada skor tablosuna bakar.",
    )
    add_image(doc, charts["competitor"])
    add_heading_styled(doc, "22.2 Personel destek masası", 2)
    body(
        doc,
        "Bildirim düşer → Ops veya liste ekranından talebi açar → üstlenir → gerekirse iç not yazar → yarışmacıya yanıtlar → "
        "durumu Bekliyor veya Çözüldü yapar → çoklu takımı etkiliyorsa duyuru üretir → gün sonunda raporlara bakar.",
    )
    add_heading_styled(doc, "22.3 Organizatör yarışma açılışı", 2)
    body(
        doc,
        "Sunucu oluşturur → rolleri kurar → kategori/kanal ağacını örer → CTFd’den takım ve görev CSV’sini alır → "
        "CTFd API bağlar → talep öneki ve politikayı ayarlar → davetleri dağıtır → canlıda skor ve talepleri izler.",
    )

    # 23
    add_heading_styled(doc, "23. Bilerek Sonraki Aşamaya Bırakılanlar", 1)
    bullets(
        doc,
        [
            "Sesli kanallar — VOICE_PLACEHOLDER; WebRTC medya henüz yok",
            "Service Worker ile arka plan push — tarayıcı Notification var, SW push yok",
            "Tam metin mesaj arama indeksi — istemci tarafı filtre ötesi",
            "Kanal sürükle-bırak yeniden sıralama mutasyonu",
            "Ayrıntılı kullanıcı bildirim tercihleri sayfasının derinleştirilmesi",
        ],
    )

    # 24
    add_heading_styled(doc, "24. Sözlük", 1)
    table_kv(
        doc,
        ["Terim", "Açıklama"],
        [
            ["CTF", "Capture The Flag — siber güvenlik yarışması formatı"],
            ["CTFd", "Yaygın CTF platformu; skor ve görev kaynağı"],
            ["Sunucu", "Bir yarışma / operasyon çalışma alanı"],
            ["Kanal", "Mesajların aktığı sohbet odası"],
            ["Ticket", "Destek talebi kaydı"],
            ["İç not", "Yalnızca personelin gördüğü ticket yorumu"],
            ["RBAC", "Role Based Access Control — role dayalı erişim"],
            ["Presence", "Çevrimiçi / boşta / rahatsız etme / çevrimdışı durumu"],
            ["First blood", "Bir görevi ilk çözen takım"],
            ["Socket.IO", "Gerçek zamanlı çift yönlü olay katmanı"],
            ["CSRF", "Sahte çapraz site isteğine karşı token koruması"],
            ["Audit log", "Hassas yönetim eylemlerinin iz kaydı"],
        ],
    )

    # Closing
    doc.add_paragraph()
    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = end.add_run(
        "— SYDT Chat · CTF Operations Console —\n"
        "Bu doküman, kod tabanındaki ürün davranışına dayalı ürün tanıtım kılavuzudur.\n"
        "İç kullanım / tescilli; yeniden dağıtım için lisans koşullarına bakınız."
    )
    set_run_font(r, 10, color=RGBColor(0x64, 0x74, 0x8B))

    doc.save(OUT)
    return OUT


def main():
    charts = {
        "architecture": chart_architecture(),
        "auth": chart_auth(),
        "competitor": chart_competitor(),
        "tickets": chart_tickets(),
        "admin": chart_admin(),
        "messaging": chart_messaging(),
        "roles": chart_roles(),
    }
    path = build_document(charts)
    print(f"Wrote {path}")
    print(f"Charts in {ASSETS}")


if __name__ == "__main__":
    main()
